"""把 docs/工程总结.md 转成排版良好的 Word 文档。

不是简单地把文本倒进去，而是解析 Markdown 结构：
  # / ## / ###  ->  Word 标题样式（可用导航窗格跳转、可自动生成目录）
  Markdown 表格 ->  真实 Word 表格（带表头底纹与边框）
  围栏代码块    ->  等宽字体段落
  引用块        ->  缩进 + 左侧强调
  - / 1. 列表   ->  项目符号 / 编号

中文字体同时设置 ascii/hAnsi 与 eastAsia，否则中文会回退成默认字体。

用法: python tools/md_to_docx.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

SRC = Path(r"D:\手机记账app\docs\工程总结.md")
DST = Path(r"D:\手机记账app\docs\工程总结.docx")

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_run_font(run, latin=CN_FONT, east=CN_FONT, size=None, bold=None, color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

# 零宽空格：给不可断的长字符串（URL、SHA256）插入换行机会。
# 不这样做的话，Word 会为了容纳这些长 token 把列撑宽，导致表格溢出页面被裁切。
ZWSP = "\u200b"


def breakable(text):
    """在 / - _ 之后以及 . 周围插入零宽断点，让长 token 能折行。"""
    for ch in ("/", "-", "_"):
        text = text.replace(ch, ch + ZWSP)
    return text.replace(".", ZWSP + "." + ZWSP)


def add_inline(par, text, size=10.5, base_bold=False, in_table=False):
    """处理 **粗体** 与 `代码` 行内标记，并给长 token 加折行点。

    in_table=True 时把等宽字体再调小一档：64 字符的 SHA256 在表格列宽里
    放不下会导致整行溢出，小一档正好能容纳。
    """
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(breakable(part[2:-2]))
            set_run_font(r, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            code_size = size - 2.0 if in_table else size - 0.5
            r = par.add_run(breakable(part[1:-1]))
            set_run_font(r, latin=MONO_FONT, east=CN_FONT, size=code_size,
                         color=(0xA0, 0x30, 0x30))
        else:
            r = par.add_run(breakable(part))
            set_run_font(r, size=size, bold=base_bold)


def set_table_width(table, total_cm, first_col_cm=None):
    """固定表格宽度并**同时**写入 tblGrid。

    关键：tblLayout=fixed 时 Word / LibreOffice 以 tblGrid 的列宽为准，
    只设单元格宽度是没用的 —— python-docx 新建表格时 tblGrid 是等宽的，
    于是两列表格会被撑成各占一半，总宽超出页面，右侧内容被裁掉。
    """
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    ncols = len(table.columns)
    if ncols == 0:
        return
    if first_col_cm is None:
        if ncols == 2:
            # 两列：第一列是字段名，占约 1/3
            first_col_cm = total_cm * 0.34
        else:
            # 多列：第一列通常是序号/短标签，给个够用的固定窄宽，
            # 否则会白白占掉 1/3，把后面的内容列挤到裁切
            first_col_cm = min(1.3, total_cm * 0.12)
    rest = total_cm - first_col_cm
    rest_each = rest / (ncols - 1) if ncols > 1 else 0

    def tw(cm):
        return int(round(cm * 567))   # 1cm = 567 twips

    widths_tw = [tw(first_col_cm)] + [tw(rest_each)] * (ncols - 1)

    # 1) 表格总宽
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(widths_tw)))

    # 2) 改写 tblGrid（决定性的一步）
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_tw:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    tblPr.addnext(grid)

    # 3) 每个单元格也写上，保持一致
    for row in tbl.findall(qn("w:tr")):
        for idx, tc in enumerate(row.findall(qn("w:tc"))):
            if idx >= len(widths_tw):
                continue
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = OxmlElement("w:tcPr")
                tc.insert(0, tcPr)
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:type"), "dxa")
            tcW.set(qn("w:w"), str(widths_tw[idx]))


def set_cell_padding(table, left_right_tw=100, top_bottom_tw=40):
    """给单元格加内边距，否则文字紧贴表格线，很难看。"""
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("w:top", top_bottom_tw), ("w:left", left_right_tw),
                     ("w:bottom", top_bottom_tw), ("w:right", left_right_tw)):
        el = OxmlElement(tag)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def build():
    lines = SRC.read_text(encoding="utf-8").split("\n")
    doc = Document()

    # 显式设成 A4 并留出页边距。
    # 注意：python-docx 默认给的是 Letter（21.59cm），而 LibreOffice 渲染时按 A4
    # 处理，两者不一致会让表格右侧被裁掉。这里统一成 A4。
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)

    # 可用宽度，再留 0.3cm 余量，避免边界处被裁
    content_cm = 21.0 - 2.2 - 2.2 - 0.3

    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    for name, sz in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5)):
        st = doc.styles[name]
        st.font.name = CN_FONT
        st.font.size = Pt(sz)
        st.font.color.rgb = RGBColor(0x1F, 0x33, 0x46)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # ---- 围栏代码块 ----
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("\n".join(buf))
            set_run_font(r, latin=MONO_FONT, east=CN_FONT, size=9)
            continue

        # ---- 表格 ----
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = ""
                add_inline(cell.paragraphs[0], h, size=10, base_bold=True, in_table=True)
                shade(cell, "EEF2F6")
            for row in rows:
                cells = table.add_row().cells
                for j, val in enumerate(row[: len(header)]):
                    cells[j].text = ""
                    add_inline(cells[j].paragraphs[0], val, size=10, in_table=True)
            set_table_width(table, content_cm)
            set_cell_padding(table)
            doc.add_paragraph()
            continue

        # ---- 标题 ----
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            h = doc.add_heading(level=level)
            add_inline(h, m.group(2), size={1: 20, 2: 15, 3: 12.5, 4: 11}[level])
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x33, 0x46)
            i += 1
            continue

        # ---- 分隔线 ----
        if re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "C9D4DE")
            bd.append(bottom)
            ppr.append(bd)
            i += 1
            continue

        # ---- 引用 ----
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            ppr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "D9A441")
            bd.append(left)
            ppr.append(bd)
            add_inline(p, " ".join(x for x in buf if x), size=10.5)
            continue

        # ---- 列表 ----
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(3))
            i += 1
            continue

        # ---- 空行 / 正文 ----
        if not stripped:
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.35
        add_inline(p, stripped)
        i += 1

    doc.save(DST)
    print(f"已生成: {DST}")
    print(f"大小  : {DST.stat().st_size} 字节")


if __name__ == "__main__":
    build()
